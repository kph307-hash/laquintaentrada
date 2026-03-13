import os
import re
import json
import math
import time
import shutil
import logging
import tempfile
import urllib.parse
from datetime import datetime
from pathlib import Path
from typing import Optional, List

import openpyxl
from dotenv import load_dotenv
from fastapi import FastAPI, Form, Request, UploadFile, File, Query, Depends, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.middleware.sessions import SessionMiddleware
from sqlalchemy import case, or_
from sqlalchemy.orm import Session
from passlib.context import CryptContext

from db import SessionLocal, init_db, Producto, UsuarioCliente, get_db, get_conn, DB_PATH

load_dotenv()

app = FastAPI(title="Super Web")
pwd_context = CryptContext(schemes=["pbkdf2_sha256"], deprecated="auto")

# ======================
# CONFIG
# ======================
PROMOS_DISK_DIR = Path(os.getenv("PROMOS_DISK_DIR", "/var/data/promos"))
PROMOS_DISK_DIR.mkdir(parents=True, exist_ok=True)
PRODUCTS_DISK_DIR = Path(os.getenv("PRODUCTS_DISK_DIR", "/var/data/productos"))
PRODUCTS_DISK_DIR.mkdir(parents=True, exist_ok=True)
SHIPPING_PER_KM = 500  # ₡ por km
SYNC_STATUS_FILE = Path("sync_status.json")
ADMIN_SYNC_TOKEN = os.getenv("ADMIN_SYNC_TOKEN", "cambie-esto")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("super_web")

ENV = os.getenv("ENV", "dev").lower()


def must_env(name: str, default: Optional[str] = None) -> str:
    v = os.getenv(name, default)
    if ENV == "prod" and (v is None or str(v).strip() == ""):
        raise RuntimeError(f"Falta variable de entorno en PROD: {name}")
    return (v or "").strip()


ADMIN_KEY = must_env("ADMIN_KEY", "1234" if ENV != "prod" else None)
SECRET_KEY = must_env("SECRET_KEY", "dev-secret" if ENV != "prod" else None)
WHATSAPP_PHONE = must_env("WHATSAPP_PHONE", "50662154752")

HTTPS_ONLY = True if ENV == "prod" else False

# Coordenadas del supermercado (Alajuela)
SUPER_LAT = 10.010209335902667
SUPER_LNG = -84.21685918761378

# Distancia máxima permitida en km
MAX_DISTANCE_KM = 3

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STATIC_DIR = os.path.join(BASE_DIR, "static")
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")

os.makedirs(STATIC_DIR, exist_ok=True)
os.makedirs(TEMPLATES_DIR, exist_ok=True)

app.add_middleware(
    SessionMiddleware,
    secret_key=SECRET_KEY,
    same_site="lax",
    https_only=HTTPS_ONLY,
)

templates = Jinja2Templates(directory=TEMPLATES_DIR)
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


@app.on_event("startup")
def startup():
    init_db()
    logger.info("✅ DB inicializada")
    logger.info("✅ TEMPLATES_DIR: %s", TEMPLATES_DIR)
    logger.info("✅ STATIC_DIR: %s", STATIC_DIR)
    logger.info("✅ ENV=%s HTTPS_ONLY=%s", ENV, HTTPS_ONLY)

    print("DB_PATH =", DB_PATH)
    print("PROMOS_DISK_DIR =", PROMOS_DISK_DIR)

@app.get("/media/promos/{filename}")
def media_promo(filename: str):
    file_path = PROMOS_DISK_DIR / filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Imagen no encontrada")
    return FileResponse(file_path)

@app.get("/media/productos/{filename}")
def media_producto(filename: str):
    file_path = PRODUCTS_DISK_DIR / filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Imagen no encontrada")
    return FileResponse(file_path)

# ======================
# HELPERS
# ======================

def is_admin(request: Request) -> bool:
    return bool(request.session.get("is_admin"))


def require_admin(request: Request) -> Optional[RedirectResponse]:
    if not is_admin(request):
        return RedirectResponse("/admin", status_code=303)
    return None

def get_cliente_session(request: Request):
    cliente_id = request.session.get("cliente_id")
    cliente_nombre = request.session.get("cliente_nombre")
    cliente_correo = request.session.get("cliente_correo")
    cliente_telefono = request.session.get("cliente_telefono")

    if not cliente_id:
        return None

    return {
        "id": cliente_id,
        "nombre": cliente_nombre,
        "correo": cliente_correo,
        "telefono": cliente_telefono,
    }


def is_cliente_logged(request: Request) -> bool:
    return bool(request.session.get("cliente_id"))

def parse_precio_cr(value) -> int:
    """Convierte precios a entero."""
    if value is None:
        return 0

    if isinstance(value, (int, float)):
        return int(round(value))

    s = str(value).replace("\xa0", " ").strip()
    if not s:
        return 0

    s = re.sub(r"[^\d,\.]", "", s)
    if not s:
        return 0

    if "." in s and "," not in s:
        parts = s.split(".")
        if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3):
            s = s.replace(".", "")

    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s and "." not in s:
        s = s.replace(",", ".")

    try:
        return int(round(float(s)))
    except Exception:
        return 0


def distance_km(lat1, lon1, lat2, lon2):
    R = 6371  # km
    dlat = math.radians(lat2 - lat1)
    dlon = math.radians(lon2 - lon1)

    a = (
        math.sin(dlat / 2) ** 2
        + math.cos(math.radians(lat1))
        * math.cos(math.radians(lat2))
        * math.sin(dlon / 2) ** 2
    )

    c = 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))
    return R * c


def parse_float_safe(v) -> float:
    try:
        if v is None:
            return 0.0
        if isinstance(v, (int, float)):
            return float(v)
        s = str(v).strip().replace(",", ".")
        if not s:
            return 0.0
        return float(s)
    except Exception:
        return 0.0


def parse_unidad_medida(value) -> str:
    s = str(value or "").strip().lower()
    s = re.sub(r"\s+", " ", s)

    if not s:
        return "UND"

    kg_tokens = (
        "kg", "kg.", "kgs", "kgs.",
        "kilo", "kilos",
        "kilogramo", "kilogramos",
        "kilogram", "kilograms",
        "kgr", "kgrs"
    )
    if any(token in s for token in kg_tokens):
        return "KG"

    und_tokens = (
        "und", "unds", "unidad", "unidades",
        "unit", "units", "pz", "pieza", "piezas"
    )
    if any(token in s for token in und_tokens):
        return "UND"

    return "UND"

def format_qty_display(qty: float, unidad_medida: str) -> str:
    unidad_medida = (unidad_medida or "UND").upper()

    if unidad_medida == "KG":
        return f"{qty:.2f} kg"

    try:
        n = int(qty)
    except Exception:
        n = 0

    return f"{n}"

def get_producto_image_url(sku: str) -> str:
    sku = (sku or "").strip()
    if not sku:
        return ""

    for ext in (".jpg", ".jpeg", ".png", ".webp"):
        filename = f"{sku}{ext}"
        path = PRODUCTS_DISK_DIR / filename
        if path.exists():
            return f"/media/productos/{filename}"

    return ""

# ======================
# HEALTHCHECK
# ======================

@app.get("/ping")
async def ping():
    return {"ok": True, "env": ENV}


@app.get("/favicon.ico", include_in_schema=False)
def favicon():
    return RedirectResponse("/static/favicon.ico", status_code=302)


# ======================
# PROMOS: DB HELPERS
# ======================

def list_promos(only_active: bool = True):
    conn = get_conn()
    cur = conn.cursor()
    if only_active:
        cur.execute("SELECT * FROM promos WHERE activo=1 ORDER BY orden ASC, id DESC")
    else:
        cur.execute("SELECT * FROM promos ORDER BY orden ASC, id DESC")
    rows = cur.fetchall()
    conn.close()
    return rows


def get_promo(promo_id: int):
    conn = get_conn()
    cur = conn.cursor()
    cur.execute("SELECT * FROM promos WHERE id=?", (promo_id,))
    row = cur.fetchone()
    conn.close()
    return row


# ======================
# HOME
# ======================

@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    promos = list_promos(only_active=True)
    return templates.TemplateResponse(
        "home.html",
        {
            "request": request,
            "promos": promos,
            "whatsapp_phone": WHATSAPP_PHONE,
            "show_search": False,
            "q": "",
            "selected_familias": [],
            "page_size": 25,
            "cliente": get_cliente_session(request),
        },
    )


# ======================
# CARRITO
# ======================

@app.get("/carrito", response_class=HTMLResponse)
@app.get("/carrito/", response_class=HTMLResponse)
def carrito(request: Request):
    return templates.TemplateResponse(
        "carrito.html",
        {
            "request": request,
            "whatsapp_phone": WHATSAPP_PHONE,
            "show_search": False,
            "q": "",
            "selected_familias": [],
            "page_size": 25,
            "cliente": get_cliente_session(request),
        },
    )


# ======================
# TIENDA
# ======================

@app.get("/tienda", response_class=HTMLResponse)
@app.get("/tienda/", response_class=HTMLResponse)
def tienda(
    request: Request,
    q: Optional[str] = None,
    page: int = 1,
    page_size: int = 25,
    fam: Optional[List[str]] = Query(default=None),
    db: Session = Depends(get_db),
):
    page = max(1, int(page or 1))
    if page_size not in (25, 50, 100):
        page_size = 25

    col_familia = getattr(Producto, "familia", None)

    familias: List[str] = []
    if col_familia is not None:
        fam_rows = (
            db.query(col_familia)
            .filter(Producto.cantidad > 0)
            .distinct()
            .all()
        )
        for (f,) in fam_rows:
            ftxt = (str(f).strip() if f else "").strip() or "Sin categoría"
            familias.append(ftxt)
        familias = sorted(set(familias))
    else:
        familias = ["Sin categoría"]

    query = db.query(Producto).filter(Producto.cantidad > 0)

    selected_familias = [x.strip() for x in (fam or []) if (x or "").strip()]

    if q and q.strip():
        q_clean = q.strip()

        query = query.filter(
            or_(
                Producto.nombre.ilike(f"%{q_clean}%"),
                Producto.sku.ilike(f"%{q_clean}%"),
            )
        ).order_by(
            case((Producto.nombre.ilike(f"{q_clean}%"), 0), else_=1),
            case((Producto.sku.ilike(f"{q_clean}%"), 0), else_=1),
            Producto.nombre.asc(),
        )
    else:
        query = query.order_by(Producto.nombre.asc())

    if selected_familias and col_familia is not None:
        query = query.filter(col_familia.in_(selected_familias))

    total = int(query.count())
    total_pages = max(1, (total + page_size - 1) // page_size)
    if page > total_pages:
        page = total_pages

    productos_page = query.offset((page - 1) * page_size).limit(page_size).all()

    productos_list = []
    for p in productos_page:
        fam_txt = (getattr(p, "familia", "") or "").strip() or "Sin categoría"
        sku_txt = (p.sku or "").strip()

        productos_list.append(
            {
                "sku": sku_txt,
                "nombre": p.nombre,
                "precio": int(p.precio),
                "familia": fam_txt,
                "cantidad": float(getattr(p, "cantidad", 0) or 0),
                "imagen_url": get_producto_image_url(sku_txt),
                "unidad_medida": (getattr(p, "unidad_medida", "UND") or "UND").upper(),
            }
        )

    print("productos_page =", len(productos_page))
    print("productos_list =", len(productos_list))
    print("request url =", request.url)

    qs_parts = []
    if q and q.strip():
        qs_parts.append(("q", q.strip()))
    qs_parts.append(("page_size", str(page_size)))
    for f in selected_familias:
        qs_parts.append(("fam", f))

    qs_base = urllib.parse.urlencode(qs_parts, doseq=True)

    ctx = {
        "request": request,
        "productos": productos_list,
        "familias": familias,
        "selected_familias": selected_familias,
        "q": q or "",
        "page": int(page),
        "page_size": int(page_size),
        "total": int(total),
        "total_pages": int(total_pages),
        "qs_base": qs_base or "",
        "whatsapp_phone": WHATSAPP_PHONE,
        "show_search": True,
        "cliente": get_cliente_session(request),
    }

    return templates.TemplateResponse("index.html", ctx)

# ======================
# PEDIDO
# ======================

@app.post("/pedido", response_class=HTMLResponse)
def crear_pedido(
    request: Request,
    delivery_mode: str = Form("delivery"),
    nombre: str = Form(""),
    correo: str = Form(""),
    telefono: str = Form(""),
    direccion: Optional[str] = Form(None),
    maps_url: Optional[str] = Form(None),
    lat: Optional[float] = Form(None),
    lng: Optional[float] = Form(None),
    cart_json: str = Form(...),
):
    cliente = get_cliente_session(request)

    nombre = (nombre or "").strip()
    correo = (correo or "").strip().lower()
    telefono = (telefono or "").strip()
    delivery_mode = (delivery_mode or "delivery").strip().lower()

    if cliente:
        if not nombre:
            nombre = cliente.get("nombre", "").strip()
        if not correo:
            correo = cliente.get("correo", "").strip().lower()

    direccion_clean = (direccion or "").strip()
    maps_url_clean = (maps_url or "").strip()

    if not nombre:
        return HTMLResponse("Falta el nombre.", status_code=400)

    if delivery_mode == "delivery" and not direccion_clean and not ((lat is not None) and (lng is not None)):
        return HTMLResponse("Falta la dirección de envío o la ubicación GPS.", status_code=400)

    # Shipping / distancia
    shipping = 0
    dist: Optional[float] = None

    if delivery_mode == "delivery":
        has_gps = (lat is not None) and (lng is not None)

        if has_gps:
            dist = distance_km(SUPER_LAT, SUPER_LNG, float(lat), float(lng))
            shipping = max(500, int(math.ceil(dist * SHIPPING_PER_KM)))

            if dist > MAX_DISTANCE_KM:
                return HTMLResponse(
                    f"""
                    <div style="max-width:640px;margin:40px auto;background:#fff;padding:18px;border-radius:16px;
                                border:1px solid rgba(0,0,0,.12);font-family:system-ui">
                      <h2>🚫 Fuera de zona de entrega</h2>
                      <p>
                        Solo entregamos dentro de <b>{MAX_DISTANCE_KM} km</b> del supermercado.<br>
                        Tu ubicación está a <b>{dist:.2f} km</b>.
                      </p>
                      <p style="margin-top:14px;">
                        <a href="/carrito" style="display:inline-block;padding:10px 14px;border-radius:12px;
                           background:#111;color:#fff;text-decoration:none;font-weight:800;">
                          Volver al carrito
                        </a>
                      </p>
                    </div>
                    """,
                    status_code=400,
                )
        else:
            shipping = 500
    else:
        shipping = 0

    # Leer carrito
    try:
        cart = json.loads(cart_json)
        if not isinstance(cart, list):
            raise ValueError()
    except Exception:
        return HTMLResponse("Carrito inválido.", status_code=400)

    total = 0
    detalle = []

    dbs = SessionLocal()
    try:
        for item in cart:
            sku = (item.get("sku") or "").strip()
            try:
                qty = float(item.get("qty", 0) or 0)
            except Exception:
                qty = 0

            unidad_medida = str(item.get("unidad_medida") or "UND").strip().upper()

            if not sku or qty <= 0:
                continue

            sku_norm = sku.strip().upper()

            # PROMOS
            if sku_norm.startswith("PROMO:") or sku_norm.startswith("PROMO-"):
                promo_nombre = (
                    (item.get("nombre") or item.get("name") or item.get("titulo") or item.get("producto") or "")
                ).strip()

                raw_price = item.get("price")
                if raw_price is None:
                    raw_price = item.get("precio")

                promo_precio = parse_precio_cr(raw_price)

                if promo_precio <= 0:
                    promo_precio = parse_precio_cr(item.get("precio_unit"))

                if not promo_nombre or promo_precio <= 0:
                    continue

                subtotal = promo_precio * qty
                total += subtotal
                detalle.append(
                    {
                        "producto": promo_nombre,
                        "precio_unit": promo_precio,
                        "qty": qty,
                        "unidad_medida": "UND",
                        "subtotal": int(round(subtotal)),
                    }
                )
                continue

            # PRODUCTOS BD
            p = dbs.query(Producto).filter(Producto.sku == sku).first()
            if not p:
                continue

            # ✅ No permitir si no hay stock
            stock_actual = int(getattr(p, "cantidad", 0) or 0)
            unidad_real = (getattr(p, "unidad_medida", "UND") or "UND").upper()

            if stock_actual <= 0:
                continue

            precio_unit = int(p.precio)

            if unidad_real == "KG":
                # qty viene en kilos: 0.25, 0.50, etc.
                if qty <= 0:
                    continue
                subtotal = int(round(precio_unit * qty))
            else:
                qty = int(round(qty))
                if qty <= 0 or qty > stock_actual:
                    continue
                subtotal = int(round(precio_unit * qty))

            total += subtotal

            detalle.append(
                {
                    "sku": p.sku,
                    "producto": p.nombre,
                    "precio_unit": precio_unit,
                    "qty": qty,
                    "unidad_medida": unidad_real,
                    "subtotal": subtotal,
                }
            )
    finally:
        dbs.close()

    if not detalle:
        return HTMLResponse("Carrito vacío (o productos sin stock).", status_code=400)

    grand_total = total + shipping

    conn = get_conn()
    cur = conn.cursor()

    usuario_id = cliente["id"] if cliente else None

    cur.execute(
        """
        INSERT INTO pedidos (
            status, usuario_id, cliente_nombre, cliente_correo, cliente_telefono,
            delivery_mode, nombre, direccion, maps_url, lat, lng, subtotal, shipping, total
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            "pendiente",
            usuario_id,
            nombre,
            correo,
            telefono,
            delivery_mode,
            nombre,
            direccion_clean,
            maps_url_clean,
            lat,
            lng,
            int(total),
            int(shipping),
            int(grand_total),
        ),
    )
    pedido_id = cur.lastrowid

    for d in detalle:
        cur.execute(
            """
            INSERT INTO pedido_items (pedido_id, sku, producto, precio_unit, qty, unidad_medida, subtotal, es_promo)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                int(pedido_id),
                d.get("sku", "") or "",
                d["producto"],
                int(d["precio_unit"]),
                float(d["qty"]),
                d.get("unidad_medida", "UND"),
                int(d["subtotal"]),
                0,
            ),
        )

    conn.commit()
    conn.close()

    # Construir mensaje WhatsApp
    mensaje = f"🛒 Pedido de: {nombre}\n"

    if delivery_mode == "delivery":
        mensaje += "🚚 Entrega: Envío a domicilio\n"
        if direccion_clean:
            mensaje += f"📍 Dirección: {direccion_clean}\n"
        if maps_url_clean:
            mensaje += f"📌 Ubicación: {maps_url_clean}\n"
        if dist is not None:
            mensaje += f"📏 Distancia: {dist:.2f} km\n"
            mensaje += f"🚚 Envío: ₡{shipping} (₡{SHIPPING_PER_KM}/km)\n"
        else:
            mensaje += "📏 Distancia: No verificada (sin GPS)\n"
            mensaje += f"🚚 Envío: ₡{shipping} mínimo (₡{SHIPPING_PER_KM}/km)\n"
    else:
        mensaje += "🏪 Entrega: Pasar a recoger en tienda\n"

    mensaje += "\nProductos:\n"
    for d in detalle:
        unidad = d.get("unidad_medida", "UND")
        if unidad == "KG":
            mensaje += f"- {d['producto']} — {format_qty_display(d['qty'], unidad)} (₡{d['precio_unit']} por kilo aprox.): ₡{d['subtotal']} aprox.\n"
        else:
            mensaje += f"- {d['producto']} x{int(round(d['qty']))} (₡{d['precio_unit']}): ₡{d['subtotal']}\n"

    if any((d.get("unidad_medida") == "KG") for d in detalle):
        mensaje += "\nNota: Los productos por kilo son aproximados y el monto final puede variar según el peso real.\n"

    mensaje += f"\n💰 TOTAL A PAGAR: ₡{grand_total}"

    texto = urllib.parse.quote(mensaje)
    whatsapp_url = f"https://wa.me/{WHATSAPP_PHONE}?text={texto}"

    return templates.TemplateResponse(
        "resultado.html",
        {
            "request": request,
            "delivery_mode": delivery_mode,
            "nombre": nombre,
            "correo": correo,
            "telefono": telefono,
            "direccion": direccion_clean,
            "maps_url": maps_url_clean,
            "detalle": detalle,
            "total": total,
            "shipping": shipping,
            "distance_km": dist,
            "grand_total": grand_total,
            "whatsapp_url": whatsapp_url,
            "pedido_id": pedido_id,
            "cliente": cliente,
            "shipping_per_km": SHIPPING_PER_KM,
            "shipping_verified": dist is not None,
        },
    )

@app.get("/login", response_class=HTMLResponse)
def cliente_login_view(request: Request):
    return templates.TemplateResponse(
        "login.html",
        {
            "request": request,
            "error": None,
            "cliente": get_cliente_session(request),
        },
    )


@app.post("/login", response_class=HTMLResponse)
def cliente_login(
    request: Request,
    correo: str = Form(...),
    password: str = Form(...),
    db: Session = Depends(get_db),
):
    correo = (correo or "").strip().lower()
    password = (password or "").strip()

    user = db.query(UsuarioCliente).filter(UsuarioCliente.correo == correo).first()

    if not user or not pwd_context.verify(password, user.password_hash):
        return templates.TemplateResponse(
            "login.html",
            {
                "request": request,
                "error": "Correo o contraseña incorrectos.",
                "cliente": get_cliente_session(request),
            },
            status_code=400,
        )

    request.session["cliente_id"] = user.id
    request.session["cliente_nombre"] = user.nombre
    request.session["cliente_correo"] = user.correo
    request.session["cliente_telefono"] = user.telefono or ""

    return RedirectResponse("/perfil", status_code=303)


@app.get("/registro", response_class=HTMLResponse)
def cliente_registro_view(request: Request):
    return templates.TemplateResponse(
        "registro.html",
        {
            "request": request,
            "error": None,
            "cliente": get_cliente_session(request),
        },
    )


@app.post("/registro", response_class=HTMLResponse)
def cliente_registro(
    request: Request,
    nombre: str = Form(...),
    correo: str = Form(...),
    telefono: str = Form(""),
    password: str = Form(...),
    db: Session = Depends(get_db),
):
    nombre = (nombre or "").strip()
    correo = (correo or "").strip().lower()
    telefono = (telefono or "").strip()
    password = (password or "").strip()

    if not nombre or not correo or not password:
        return templates.TemplateResponse(
            "registro.html",
            {
                "request": request,
                "error": "Complete nombre, correo y contraseña.",
                "cliente": get_cliente_session(request),
            },
            status_code=400,
        )

    existe = db.query(UsuarioCliente).filter(UsuarioCliente.correo == correo).first()
    if existe:
        return templates.TemplateResponse(
            "registro.html",
            {
                "request": request,
                "error": "Ese correo ya está registrado.",
                "cliente": get_cliente_session(request),
            },
            status_code=400,
        )
    if len(password.encode("utf-8")) > 72:
        return templates.TemplateResponse(
            "registro.html",
            {
                "request": request,
                "error": "La contraseña es demasiado larga. Use una de máximo 72 bytes.",
                "cliente": get_cliente_session(request),
            },
            status_code=400,
        )
    
    nuevo = UsuarioCliente(
        nombre=nombre,
        correo=correo,
        telefono=telefono,
        password_hash=pwd_context.hash(password),
    )

    db.add(nuevo)
    db.commit()
    db.refresh(nuevo)

    request.session["cliente_id"] = nuevo.id
    request.session["cliente_nombre"] = nuevo.nombre
    request.session["cliente_correo"] = nuevo.correo
    request.session["cliente_telefono"] = nuevo.telefono or ""

    return RedirectResponse("/perfil", status_code=303)


@app.get("/logout")
def cliente_logout(request: Request):
    request.session.pop("cliente_id", None)
    request.session.pop("cliente_nombre", None)
    request.session.pop("cliente_correo", None)
    request.session.pop("cliente_telefono", None)
    return RedirectResponse("/", status_code=303)


@app.get("/perfil", response_class=HTMLResponse)
def cliente_perfil(request: Request):
    cliente = get_cliente_session(request)
    if not cliente:
        return RedirectResponse("/login", status_code=303)

    conn = get_conn()
    cur = conn.cursor()
    cur.execute(
        """
        SELECT * FROM pedidos
        WHERE usuario_id = ?
        ORDER BY id DESC
        """,
        (cliente["id"],),
    )
    pedidos = cur.fetchall()
    conn.close()

    return templates.TemplateResponse(
        "perfil.html",
        {
            "request": request,
            "cliente": cliente,
            "pedidos": pedidos,
        },
    )

# ======================
# ADMIN
# ======================

def leer_estado_sync():
    if not SYNC_STATUS_FILE.exists():
        return None

    data = json.loads(SYNC_STATUS_FILE.read_text(encoding="utf-8"))

    iso = data.get("ultima_actualizacion")
    if iso:
        dt = datetime.fromisoformat(iso)
        data["ultima_actualizacion_legible"] = dt.strftime("%d/%m/%Y %I:%M %p")

    return data


@app.get("/admin", response_class=HTMLResponse)
def admin_panel(request: Request, db: Session = Depends(get_db)):
    if not is_admin(request):
        return templates.TemplateResponse(
            "admin_login.html",
            {"request": request, "error": None},
        )

    productos = db.query(Producto).order_by(Producto.nombre.asc()).all()
    estado_sync = leer_estado_sync()

    return templates.TemplateResponse(
        "admin.html",
        {
            "request": request,
            "productos": productos,
            "estado_sync": estado_sync,
        },
    )


@app.post("/admin/login", response_class=HTMLResponse)
def admin_login(request: Request, key: str = Form(...)):
    if key == ADMIN_KEY:
        request.session["is_admin"] = True
        return RedirectResponse("/admin", status_code=303)

    return templates.TemplateResponse(
        "admin_login.html",
        {"request": request, "error": "Clave incorrecta"},
    )


@app.get("/admin/logout")
def admin_logout(request: Request):
    request.session.pop("is_admin", None)
    return RedirectResponse("/admin", status_code=303)

def procesar_sync_xlsx_desde_archivo(file_path: str, db: Session):
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active

    header_row_idx = 4
    headers = next(
        ws.iter_rows(
            min_row=header_row_idx,
            max_row=header_row_idx,
            values_only=True,
        )
    )

    def norm_header(h):
        return re.sub(r"\s+", " ", str(h or "").strip().lower())

    idx = {norm_header(h): i for i, h in enumerate(headers) if h is not None}

    print("HEADERS XLSX NORMALIZADOS:", list(idx.keys()))

    col_sku = idx.get("código")
    col_nombre = idx.get("descripción")
    col_precio = idx.get("precio de venta")
    col_familia = idx.get("familia")
    col_cantidad = idx.get("inventario")

    col_unidad = idx.get("unidades")
    if col_unidad is None:
        col_unidad = idx.get("unidad")
    if col_unidad is None:
        col_unidad = idx.get("unidad de medida")

    print("COL_UNIDAD =", col_unidad)

    if col_sku is None or col_nombre is None or col_precio is None:
        wb.close()
        raise ValueError("No encuentro columnas: 'Código', 'Descripción', 'Precio de venta'.")

    existentes = db.query(Producto).all()
    by_sku = {(p.sku or "").strip(): p for p in existentes if (p.sku or "").strip()}

    creados = 0
    actualizados = 0
    saltados = 0

    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        sku = row[col_sku] if col_sku is not None and col_sku < len(row) else None
        nombre = row[col_nombre] if col_nombre is not None and col_nombre < len(row) else None
        precio = row[col_precio] if col_precio is not None and col_precio < len(row) else None

        familia = None
        if col_familia is not None and col_familia < len(row):
            familia = row[col_familia]

        cantidad = 0.0
        if col_cantidad is not None and col_cantidad < len(row):
            cantidad = parse_float_safe(row[col_cantidad])

        raw_unidad = None
        if col_unidad is not None and col_unidad < len(row):
            raw_unidad = row[col_unidad]

        unidad_medida = parse_unidad_medida(raw_unidad)

        if not sku:
            saltados += 1
            continue

        sku_txt = str(sku).strip()
        if not sku_txt:
            saltados += 1
            continue

        nombre_txt = str(nombre).strip() if nombre else ""
        if not nombre_txt:
            saltados += 1
            continue

        if sku_txt == "013608":
            print("SKU 013608 | RAW UNIDAD =", raw_unidad, "| FINAL =", unidad_medida)

        precio_int = parse_precio_cr(precio)
        if precio_int <= 0:
            saltados += 1
            continue

        familia_txt = str(familia).strip() if familia else ""

        p = by_sku.get(sku_txt)
        if p:
            changed = False

            if (p.nombre or "") != nombre_txt:
                p.nombre = nombre_txt
                changed = True

            if int(p.precio) != precio_int:
                p.precio = precio_int
                changed = True

            if (p.familia or "") != familia_txt:
                p.familia = familia_txt
                changed = True

            cantidad_actual = float(getattr(p, "cantidad", 0) or 0)
            if abs(cantidad_actual - float(cantidad)) > 0.0001:
                p.cantidad = float(cantidad)
            changed = True

            if (getattr(p, "unidad_medida", "UND") or "UND") != unidad_medida:
                p.unidad_medida = unidad_medida
                changed = True

            if changed:
                actualizados += 1
        else:
            nuevo = Producto(
                sku=sku_txt,
                nombre=nombre_txt,
                precio=precio_int,
                familia=familia_txt,
                cantidad=int(cantidad),
                unidad_medida=unidad_medida,
            )
            db.add(nuevo)
            by_sku[sku_txt] = nuevo
            creados += 1

    db.commit()
    wb.close()

    return {
        "creados": creados,
        "actualizados": actualizados,
        "saltados": saltados,
        "total": creados + actualizados,
    }

# ======================
# ADMIN: SYNC XLSX por SKU
# ======================

@app.post("/admin/sync-xlsx", response_class=HTMLResponse)
async def sync_xlsx(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    print("=== ENTRO A /admin/sync-xlsx ===")
    denied = require_admin(request)
    if denied:
        return denied

    if not (file.filename or "").lower().endswith(".xlsx"):
        return HTMLResponse("Subí un archivo .xlsx", status_code=400)

    tmp_file_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
            tmp_file_path = tmp.name
            tmp.write(await file.read())

        resultado = procesar_sync_xlsx_desde_archivo(tmp_file_path, db)

        guardar_estado_sync(
            ok=True,
            mensaje="Sincronización completada",
            total_productos=resultado["total"],
            archivo=file.filename or "archivo.xlsx",
        )

        return HTMLResponse(
            f"<h2>Sync por SKU completado ✅</h2>"
            f"<p>Creados: {resultado['creados']}</p>"
            f"<p>Actualizados: {resultado['actualizados']}</p>"
            f"<p>Saltados: {resultado['saltados']}</p>"
            f'<p><a href="/admin">Volver al panel</a></p>'
        )

    except Exception as e:
        logger.exception("Error en sync-xlsx")
        guardar_estado_sync(
            ok=False,
            mensaje=str(e),
            total_productos=0,
            archivo=file.filename or "",
        )
        return HTMLResponse(f"Error en sync: {e}", status_code=500)

    finally:
        if tmp_file_path and os.path.exists(tmp_file_path):
            try:
                os.remove(tmp_file_path)
            except Exception:
                pass


# ======================
# ADMIN: PROMOS
# ======================

@app.get("/admin/promos", response_class=HTMLResponse, include_in_schema=False)
def admin_promos(request: Request):
    denied = require_admin(request)
    if denied:
        return denied

    promos = list_promos(only_active=False)
    return templates.TemplateResponse(
        "admin_promos.html",
        {"request": request, "promos": promos},
    )


@app.post("/admin/promos/create", include_in_schema=False)
async def admin_promos_create(
    request: Request,
    titulo: str = Form(...),
    descripcion: str = Form(""),
    precio: str = Form(""),
    tipo: str = Form("publicidad"),
    orden: int = Form(0),
    activo: Optional[str] = Form(None),
    imagen: Optional[UploadFile] = File(None),
):
    denied = require_admin(request)
    if denied:
        return denied

    activo_int = 1 if activo else 0
    imagen_url = ""

    if imagen and imagen.filename:
        ext = os.path.splitext(imagen.filename)[1].lower() or ".jpg"
        safe_name = f"promo_{int(time.time())}{ext}"
        path = PROMOS_DISK_DIR / safe_name

        content = await imagen.read()
        with open(path, "wb") as f:
            f.write(content)

        imagen_url = f"/media/promos/{safe_name}"

    conn = get_conn()
    cur = conn.cursor()
    cur.execute(
        """
        INSERT INTO promos (titulo, descripcion, precio, tipo, imagen_url, activo, orden)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (titulo, descripcion, precio, tipo, imagen_url, activo_int, orden),
    )
    conn.commit()
    conn.close()

    return RedirectResponse("/admin/promos", status_code=303)


@app.post("/admin/promos/{promo_id}/update", include_in_schema=False)
async def admin_promos_update(
    promo_id: int,
    request: Request,
    titulo: str = Form(...),
    descripcion: str = Form(""),
    precio: str = Form(""),
    tipo: str = Form("publicidad"),
    orden: int = Form(0),
    activo: Optional[str] = Form(None),
    imagen: Optional[UploadFile] = File(None),
):
    denied = require_admin(request)
    if denied:
        return denied

    promo = get_promo(promo_id)
    if not promo:
        return RedirectResponse("/admin/promos", status_code=303)

    activo_int = 1 if activo else 0
    imagen_url = promo["imagen_url"] or ""

    if imagen and imagen.filename:
        ext = os.path.splitext(imagen.filename)[1].lower() or ".jpg"
        safe_name = f"promo_{promo_id}_{int(time.time())}{ext}"
        path = PROMOS_DISK_DIR / safe_name

        content = await imagen.read()
        with open(path, "wb") as f:
            f.write(content)

        imagen_url = f"/media/promos/{safe_name}"

    conn = get_conn()
    cur = conn.cursor()
    cur.execute(
        """
        UPDATE promos
        SET titulo=?, descripcion=?, precio=?, tipo=?, imagen_url=?, activo=?, orden=?
        WHERE id=?
        """,
        (titulo, descripcion, precio, tipo, imagen_url, activo_int, orden, promo_id),
    )
    conn.commit()
    conn.close()

    return RedirectResponse("/admin/promos", status_code=303)


@app.post("/admin/promos/{promo_id}/delete", include_in_schema=False)
def admin_promos_delete(promo_id: int, request: Request):
    denied = require_admin(request)
    if denied:
        return denied

    conn = get_conn()
    cur = conn.cursor()
    cur.execute("DELETE FROM promos WHERE id=?", (promo_id,))
    conn.commit()
    conn.close()

    return RedirectResponse("/admin/promos", status_code=303)

# ======================
# ADMIN: PEDIDOS
# ======================

@app.get("/admin/pedidos", response_class=HTMLResponse)
def admin_pedidos(request: Request):
    denied = require_admin(request)
    if denied:
        return denied

    conn = get_conn()
    cur = conn.cursor()
    cur.execute("""
      SELECT * FROM pedidos
      WHERE status='pendiente'
      ORDER BY id DESC
    """)
    pedidos = cur.fetchall()
    conn.close()

    return templates.TemplateResponse(
        "admin_pedidos.html",
        {"request": request, "pedidos": pedidos},
    )


@app.get("/admin/pedidos/archivados", response_class=HTMLResponse)
def admin_pedidos_archivados(request: Request):
    denied = require_admin(request)
    if denied:
        return denied

    conn = get_conn()
    cur = conn.cursor()
    cur.execute("""
      SELECT * FROM pedidos
      WHERE status='entregado'
      ORDER BY entregado_en DESC, id DESC
    """)
    pedidos = cur.fetchall()
    conn.close()

    return templates.TemplateResponse(
        "admin_pedidos_archivados.html",
        {"request": request, "pedidos": pedidos},
    )


@app.get("/admin/pedidos/{pedido_id}", response_class=HTMLResponse)
def admin_pedido_detalle(pedido_id: int, request: Request):
    denied = require_admin(request)
    if denied:
        return denied

    conn = get_conn()
    cur = conn.cursor()

    cur.execute("SELECT * FROM pedidos WHERE id=?", (pedido_id,))
    pedido = cur.fetchone()

    cur.execute("SELECT * FROM pedido_items WHERE pedido_id=? ORDER BY id ASC", (pedido_id,))
    items = cur.fetchall()

    conn.close()

    if not pedido:
        return HTMLResponse("Pedido no encontrado.", status_code=404)

    return templates.TemplateResponse(
        "admin_pedido_detalle.html",
        {"request": request, "pedido": pedido, "items": items},
    )


@app.post("/admin/pedidos/{pedido_id}/entregado")
def admin_pedido_entregado(pedido_id: int, request: Request):
    denied = require_admin(request)
    if denied:
        return denied

    conn = get_conn()
    cur = conn.cursor()
    cur.execute("""
      UPDATE pedidos
      SET status='entregado', entregado_en=datetime('now')
      WHERE id=?
    """, (pedido_id,))
    conn.commit()
    conn.close()

    return RedirectResponse("/admin/pedidos", status_code=303)



# ======================
# Atajo opcional a tienda
# ======================

@app.get("/go-tienda")
def go_tienda():
    return RedirectResponse("/tienda", status_code=302)

@app.get("/test-login-route")
def test_login_route():
    return {"ok": True}

def guardar_estado_sync(ok: bool, mensaje: str, total_productos: int = 0, archivo: str = ""):
    data = {
        "ok": ok,
        "mensaje": mensaje,
        "total_productos": total_productos,
        "archivo": archivo,
        "ultima_actualizacion": datetime.now().isoformat(),
    }
    SYNC_STATUS_FILE.write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8"
    )


@app.post("/admin/sync-upload")
async def admin_sync_upload(
    token: str = Form(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    print("=== ENTRO A /admin/sync-upload ===")

    if token != ADMIN_SYNC_TOKEN:
        raise HTTPException(status_code=403, detail="Token inválido")

    filename = (file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Solo se permiten archivos .xlsx")

    tmp_dir = Path("tmp")
    tmp_dir.mkdir(exist_ok=True)
    tmp_file = tmp_dir / (file.filename or "archivo.xlsx")

    with tmp_file.open("wb") as f:
        shutil.copyfileobj(file.file, f)

    try:
        resultado = procesar_sync_xlsx_desde_archivo(str(tmp_file), db)

        guardar_estado_sync(
            ok=True,
            mensaje="Sincronización completada",
            total_productos=resultado["total"],
            archivo=file.filename or "archivo.xlsx",
        )

        return {
            "ok": True,
            "mensaje": "Sincronización completada",
            "creados": resultado["creados"],
            "actualizados": resultado["actualizados"],
            "saltados": resultado["saltados"],
            "total_productos": resultado["total"],
            "archivo": file.filename,
        }

    except Exception as e:
        logger.exception("Error en sync-upload")
        guardar_estado_sync(
            ok=False,
            mensaje=str(e),
            total_productos=0,
            archivo=file.filename or "",
        )
        raise HTTPException(status_code=500, detail=f"Error sincronizando: {e}")

    finally:
        if tmp_file.exists():
            tmp_file.unlink()

@app.get("/admin/productos-imagenes", response_class=HTMLResponse)
def admin_productos_imagenes(request: Request):
    denied = require_admin(request)
    if denied:
        return denied

    return templates.TemplateResponse(
        "admin_productos_imagenes.html",
        {
            "request": request,
            "resultado": None,
        },
    )

@app.post("/admin/productos-imagenes", response_class=HTMLResponse)
async def admin_productos_imagenes_upload(
    request: Request,
    files: list[UploadFile] = File(...),
    db: Session = Depends(get_db),
):
    denied = require_admin(request)
    if denied:
        return denied

    subidas = []
    no_encontrados = []
    errores = []

    permitidas = {".jpg", ".jpeg", ".png", ".webp"}

    for f in files:
        try:
            filename = (f.filename or "").strip()
            if not filename:
                continue

            ext = os.path.splitext(filename)[1].lower()
            sku = os.path.splitext(filename)[0].strip()

            if ext not in permitidas:
                errores.append(f"{filename}: formato no permitido")
                continue

            if not sku:
                errores.append(f"{filename}: nombre inválido")
                continue

            producto = db.query(Producto).filter(Producto.sku == sku).first()
            if not producto:
                no_encontrados.append(filename)
                continue

            # borrar archivos viejos del mismo SKU con otras extensiones
            for old_ext in (".jpg", ".jpeg", ".png", ".webp"):
                old_path = PRODUCTS_DISK_DIR / f"{sku}{old_ext}"
                if old_path.exists():
                    try:
                        old_path.unlink()
                    except Exception:
                        pass

            save_path = PRODUCTS_DISK_DIR / f"{sku}{ext}"

            content = await f.read()
            with open(save_path, "wb") as out:
                out.write(content)

            subidas.append(filename)

        except Exception as e:
            errores.append(f"{getattr(f, 'filename', 'archivo')}: {e}")

    resultado = {
        "subidas": subidas,
        "no_encontrados": no_encontrados,
        "errores": errores,
        "total_subidas": len(subidas),
        "total_no_encontrados": len(no_encontrados),
        "total_errores": len(errores),
    }

    return templates.TemplateResponse(
        "admin_productos_imagenes.html",
        {
            "request": request,
            "resultado": resultado,
        },
    )


@app.get("/debug-xlsx")
def debug_xlsx():
    import tempfile
    import requests
    import openpyxl

    url = os.getenv("SHEET_XLSX_URL", "").strip()
    if not url:
        return {"ok": False, "error": "Falta SHEET_XLSX_URL en .env"}

    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        r = requests.get(url, timeout=30)
        r.raise_for_status()
        tmp.write(r.content)
        tmp_path = tmp.name

    wb = openpyxl.load_workbook(tmp_path, read_only=True, data_only=True)
    ws = wb.active

    header_row_idx = 4
    headers = next(ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx, values_only=True))
    idx = {str(h).strip(): i for i, h in enumerate(headers) if h is not None}

    col_sku = idx.get("Código")
    col_unidad = idx.get("Unidades")

    resultado = {
        "headers": list(idx.keys()),
        "col_unidad": col_unidad,
        "muestra": []
    }

    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        sku = row[col_sku] if col_sku is not None and col_sku < len(row) else None
        unidad = row[col_unidad] if col_unidad is not None and col_unidad < len(row) else None

        sku_txt = str(sku).strip() if sku else ""
        if sku_txt in ("013608", "013609", "013610"):
            resultado["muestra"].append({
                "sku": sku_txt,
                "raw_unidad": unidad,
                "final_unidad": parse_unidad_medida(unidad),
            })

    wb.close()
    os.remove(tmp_path)
    return resultado